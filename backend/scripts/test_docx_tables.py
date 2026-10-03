"""3.docx 表格结构体检：python-docx 到底能不能看到表格、段落长什么样。

背景（ADR 0009 验证关卡衍生排查）：PDF 渲染显示每页是一张带边框的课表表格，
但之前结论是"文本流"。本脚本打印段落数/表格数/表格行列数与首表样例，
判断是否可以利用表格结构修复周次归属，而不必动用 OCR。
"""

import io
import sys
from pathlib import Path


def main() -> None:
    data = Path(sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\wings\Downloads\3.docx").read_bytes()
    from docx import Document

    doc = Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    print(f"非空段落数: {len(paras)}")
    print(f"表格数: {len(doc.tables)}")
    for t_idx, table in enumerate(doc.tables):
        print(f"\n=== 表 {t_idx + 1}: {len(table.rows)} 行 x {len(table.columns)} 列 ===")
        for r_idx, row in enumerate(table.rows[:12]):
            cells = [c.text.strip().replace("\n", "/") for c in row.cells]
            print(f"  行{r_idx}: {cells}")
    print("\n--- 前 10 个非空段落 ---")
    for p in paras[:10]:
        print(f"  | {p[:80]}")


if __name__ == "__main__":
    main()
